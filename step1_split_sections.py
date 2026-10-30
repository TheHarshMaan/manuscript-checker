"""
Manuscript Checker - Milestone 1
Reads a .docx manuscript and splits it into sections.

How to run (in a terminal, from this folder):
    python step1_split_sections.py my_manuscript.docx
"""

import sys
import re
from docx import Document

# Section names we look for. Each key is the "official" name;
# the list holds other ways authors write it.
SECTION_NAMES = {
    "Abstract":              ["abstract", "summary", "structured abstract"],
    "Keywords":              ["keywords", "key words", "index terms", "key terms"],
    "Highlights":            ["highlights", "key points", "key messages", "what is already known",
                              "what this study adds", "research highlights"],
    "Introduction":          ["introduction", "background"],
    "Methods":               ["methods", "materials and methods", "material and methods", "methodology",
                              "patients and methods", "methods and materials", "subjects and methods",
                              "materials & methods", "study design and methods", "methods and analysis",
                              "method", "experimental", "experimental section", "experimental procedures",
                              "participants and methods", "proposed method", "proposed approach",
                              "materials and experimental methods", "research methodology", "methods and design"],
    "Related work":          ["related work", "related works", "literature review", "review of literature",
                              "theoretical framework", "theoretical background"],
    "Results":               ["results", "findings", "experiments", "experimental results", "evaluation",
                              "results and analysis", "experiments and results"],
    "Results and discussion": ["results and discussion", "results & discussion"],
    "Case presentation":     ["case presentation", "case report", "case description", "case history",
                              "case details", "case summary"],
    "Discussion":            ["discussion", "comment"],
    "Conclusion":            ["conclusion", "conclusions", "concluding remarks"],
    "Acknowledgements":      ["acknowledgements", "acknowledgments", "acknowledgement", "acknowledgment"],
    "Funding":               ["funding", "funding statement", "financial support", "funding sources",
                              "funding information", "source of funding", "sources of funding"],
    "Conflict of interest":  ["conflict of interest", "conflicts of interest", "competing interests",
                              "competing interest", "declaration of interest", "declaration of interests",
                              "declaration of competing interest", "declaration of competing interests",
                              "disclosure", "disclosures"],
    "Author contributions":  ["author contributions", "authors' contributions", "author contribution",
                              "contributorship", "credit authorship contribution statement",
                              "credit author statement", "authorship contributions"],
    "Ethics":                ["ethics", "ethical approval", "ethics statement", "ethical considerations",
                              "ethics approval and consent to participate", "ethics approval", "ethical statement"],
    "Consent":               ["consent", "consent for publication", "patient consent", "informed consent"],
    "Data availability":     ["data availability", "data availability statement", "availability of data and materials",
                              "data sharing statement"],
    "Data records":          ["data records", "technical validation", "usage notes", "data description"],
    "References":            ["references", "bibliography", "literature cited", "reference list", "works cited",
                              "cited works", "sources"],
}


def clean(text):
    """Lowercase, strip numbering like '2.' or 'II.' and trailing colons."""
    text = text.strip().lower()
    text = re.sub(r"^([0-9]+(\.[0-9]+)*|[ivx]+)[.)]?\s+", "", text)  # remove "2." / "II."
    return text.rstrip(":. ").strip()


def match_section(paragraph):
    """Return the official section name if this paragraph is a heading, else None."""
    text = paragraph.text.strip()
    if not text:
        return None

    # "Keywords: x, y, z" written on one line counts as the Keywords section
    if re.match(r"^key\s?words\s*[:\-]", text, re.IGNORECASE):
        return "Keywords"

    if len(text) > 60:                      # real headings are short
        return None

    cleaned = clean(text)
    for official, variants in SECTION_NAMES.items():
        if cleaned in variants:
            return official
    return None


UNITS = re.compile(r"(?:\b(?:m|cm|mm|km|kg|µm|μm|nm|ft|in|s)|\d)$")


def paragraph_text(p):
    """The paragraph's text, with superscript citation numbers written as [1,2],
    so 'diabetes triples the risk.¹,²' reads like 'diabetes triples the risk.[1,2]'.
    Superscripts after a number or a unit (10⁵, m², cm³) are left alone."""
    out = ""
    for run in p.runs:
        text = run.text
        if run.font.superscript and re.fullmatch(r"\s*\d+(?:\s*[,–-]\s*\d+)*\s*", text or "") \
                and out and not UNITS.search(out.rstrip()):
            out += "[" + re.sub(r"\s+", "", text) + "]"
        else:
            out += text
    return out if p.runs else p.text


def open_document(path, name=None):
    """A Word document, or a stand-in for PDF, ODT, RTF and LaTeX files (see readers.py)."""
    from readers import READERS, file_kind
    kind = file_kind(path, name)
    if kind == ".doc":
        raise ValueError("old Word .doc files can't be read; in Word, use File → Save As → .docx")
    if kind in READERS:
        return READERS[kind](path)
    return Document(path)


def split_into_sections(path, name=None):
    """Read the manuscript and return a list of (section_name, [paragraph texts]), and the document."""
    doc = open_document(path, name)
    sections = [("Front matter", [])]        # title, authors etc. before the first heading

    for p in doc.paragraphs:
        name = match_section(p)
        text = paragraph_text(p).strip()
        if name:
            sections.append((name, []))
            # keep inline text, e.g. "Keywords: x, y" -> "x, y"
            rest = re.split(r"[:\-]", p.text, maxsplit=1)
            if name == "Keywords" and len(rest) == 2 and rest[1].strip():
                sections[-1][1].append(rest[1].strip())
        elif text:
            sections[-1][1].append(text)

    return sections, doc


def main():
    if len(sys.argv) < 2:
        print("Usage: python step1_split_sections.py <manuscript.docx>")
        return

    path = sys.argv[1]
    try:
        sections, doc = split_into_sections(path)
    except Exception as e:
        print(f"Could not read '{path}': {e}")
        return

    print(f"\nFile: {path}")
    print(f"Tables in document: {len(doc.tables)}\n")
    print(f"{'Section':<24}{'Paragraphs':>11}{'Words':>8}")
    print("-" * 43)
    for name, paras in sections:
        words = sum(len(t.split()) for t in paras)
        print(f"{name:<24}{len(paras):>11}{words:>8}")

    found = {name for name, _ in sections}
    missing = [s for s in ["Abstract", "Keywords", "Introduction", "Methods",
                           "Results", "Discussion", "References"] if s not in found]
    print()
    if missing:
        print("Not found: " + ", ".join(missing))
    else:
        print("All core sections found.")


if __name__ == "__main__":
    main()
