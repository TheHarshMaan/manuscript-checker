"""
Reference integrity checker, by The Grey Swan Editorial.

Checks each reference in a manuscript (or a pasted reference list) against Crossref:
  - the DOI exists, and belongs to this reference (title, first author, year)
  - a reference without a DOI can be found, and its DOI is suggested
  - a reference that can't be found anywhere is flagged (it may be wrong, or made up)
  - retracted papers, expressions of concern and corrections, from Retraction Watch and Crossref

The checking logic is plain Python with no internet access of its own: lookups are passed in, so the
same code runs on the command line (urllib), on the website (the browser fetches) and in the tests
(saved answers).

How to run (in a terminal, from this folder):
    python refcheck.py manuscript.docx                 -> checks the reference list, saves reference_report.xlsx
    python refcheck.py references.txt                  -> a plain list of references, one per line
    python refcheck.py manuscript.docx --word          -> also saves reference_report.docx
"""

import gzip
import io
import json
import re
import sys
import time
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
from difflib import SequenceMatcher
from pathlib import Path

I = re.IGNORECASE
CONTACT = "editorial@thegreyswan.in"        # sent to Crossref so it can get in touch about our requests
API = "https://api.crossref.org/works"
RETRACTIONS_URL = "https://raw.githubusercontent.com/TheHarshMaan/manuscript-checker/retraction-data/retractions.json.gz"
SEARCH_FIELDS = "DOI,title,author,issued,published-print,published-online,container-title,type"

# ---------------------------------------------------------------------------
# What each result means: label, severity, how to fix
# ---------------------------------------------------------------------------
STATUSES = {
    "retracted":     ("Retracted", "must",
                      "Remove the citation, or cite it only when discussing the retraction, saying it was retracted "
                      "and citing the retraction notice (COPE retraction guidelines)."),
    "concern":       ("Expression of concern", "should",
                      "Read the notice. If you keep the citation, mention the expression of concern."),
    "doi_not_found": ("DOI not found", "should",
                      "Check the DOI for typos: copy it from the article's page or from doi.org."),
    "mismatch":      ("DOI points to a different paper", "should",
                      "Replace the DOI with the one for this reference, or correct the reference."),
    "not_found":     ("Couldn't be found", "should",
                      "Check that the reference exists and its details are right. If it's a real source that isn't in "
                      "Crossref, you can ignore this."),
    "corrected":     ("Corrected", "advisory",
                      "Check the correction doesn't change the finding you cite; cite the correction too if it does."),
    "reinstated":    ("Reinstated after retraction", "advisory",
                      "The paper was retracted and later reinstated: check the notices."),
    "details":       ("Details differ", "advisory",
                      "Check the title, first author and year against the published article."),
    "doi_missing":   ("DOI missing", "advisory", "Add the DOI to the reference."),
    "not_checked":   ("Not checked", "info",
                      "Books, reports, guidelines and websites usually aren't in Crossref: check these by hand."),
    "error":         ("Couldn't check", "info", "Crossref didn't answer; try again later."),
    "verified":      ("Verified", "ok", ""),
}
SEVERITY_ORDER = {"must": 0, "should": 1, "advisory": 2, "info": 3, "ok": 4}
SEVERITY_LABEL = {"must": "Must fix", "should": "Should fix", "advisory": "Advisory", "info": "Not checked",
                  "ok": "Verified"}

# ---------------------------------------------------------------------------
# Reading references
# ---------------------------------------------------------------------------
DOI_RE = re.compile(r"\b(10\.\d{4,9}/[^\s\"<>]+)", I)
NUMBER = re.compile(r"^\s*\[?\d+[.)\]]?\s+")
NOT_JOURNAL = re.compile(
    r"\b(?:press|publishers?|publishing|isbn|edition|eds?\.|editors?|retrieved|accessed|available (?:at|from)|"
    r"report|guidelines?|world health organi[sz]ation|ministry|government|department of health|thesis|dissertation|"
    r"handbook|manual|textbook|chapter|in:|website|web page|blog|news|act,? \d{4}|policy|white paper|preprint)\b"
    r"|https?://(?!(?:dx\.)?doi\.org)", I)


def find_doi(text):
    """The DOI in a reference, however it is written ('doi:10…', 'https://doi.org/10…'), in lower case."""
    text = re.sub(r"(?<=10\.\d{4})\s*/\s*", "/", text)                # "10.1016 / j.x" -> "10.1016/j.x"
    m = DOI_RE.search(text)
    if not m:
        return None
    doi = m.group(1).rstrip(".,;:)]}>'\"’”")
    if doi.count("(") > doi.count(")"):                                 # keep balanced brackets: S0140-6736(97)…
        doi = doi.rstrip("(")
    return doi.lower()


def strip_number(text):
    return NUMBER.sub("", text, count=1).strip()


def title_guess(reference):
    """The title as written in the reference, for comparing with Crossref (None if it can't be told)."""
    text = strip_number(reference)
    quoted = re.search(r"[“\"]([^”\"]{10,300})[”\"]", text)                               # IEEE: “Title,”
    if quoted and len(quoted.group(1).split()) >= 3:
        return quoted.group(1).strip(" ,")
    apa = re.search(r"\((?:19|20)\d{2}[a-z]?[^)]{0,20}\)\.?\s+(.+?)[.?!](?:\s|$)", text)        # APA: (2020). Title.
    if apa and len(apa.group(1).split()) >= 3:
        return apa.group(1).strip(" ,")
    parts = re.split(r"\.\s+", text)       # Vancouver: "Davey P, Marwick CA. Title. Journal. 2017;…"
    if len(parts) >= 2 and len(parts[1].split()) >= 3 and not re.match(r"(?:19|20)\d{2}", parts[1]):
        return parts[1].strip()
    chicago = re.search(r"\b(?:19|20)\d{2}[a-z]?\.\s+(.+?)[.?!](?:\s|$)", text)              # Chicago: 2020. Title.
    if chicago and len(chicago.group(1).split()) >= 3:
        return chicago.group(1).strip(" ,")
    return None


JOURNAL_SIGNS = re.compile(r"\d+\s*\(\d+\)|;\s*\d+|\bvol\.|\b\d+\s*[:,]\s*e?\d+\s*[-–]\s*\d+|\be\d{4,}\b|\bpp?\.\s*\d", I)
BOOK_PLACE = re.compile(r"\b[A-Z][a-z]+(?:[ ,]+[A-Z][a-z]+)*:\s*[A-Z][\w&.'’ ]{2,}")    # "New York: Pantheon"


def reference_kind(reference, doi):
    """'journal' for anything Crossref should know (articles, and anything with a DOI), else 'other'
    (books, reports, guidelines, websites), which aren't flagged when Crossref can't find them."""
    if doi:
        return "journal"
    if NOT_JOURNAL.search(reference) or (not JOURNAL_SIGNS.search(reference) and BOOK_PLACE.search(reference)):
        return "other"
    return "journal"


def reference_years(reference):
    return [int(y) for y in re.findall(r"(?<![\d/.])((?:19|20)\d{2})(?![\d/])", DOI_RE.sub(" ", reference))]


def split_reference_text(text):
    """A pasted reference list -> one string per reference (wrapped lines joined back together)."""
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    if not lines:
        return []
    numbered = sum(1 for l in lines if NUMBER.match(l)) >= len(lines) / 2
    refs = []
    for line in lines:
        starts = NUMBER.match(line) if numbered else (
            re.match(r"^[A-Z][\w'’\-]+(?: [A-Z][\w'’\-]+)*,? (?:[A-Z]\.|[A-Z]{1,3}\b|[A-Z][a-z]+)", line)
            or re.match(r"^[A-Z][\w &'’\-]+\.\s*\((?:19|20)\d{2}", line))
        if starts or not refs:
            refs.append(line)
        else:
            refs[-1] += (" " if not refs[-1].endswith("-") else "") + line
    return refs


def extract(path=None, name=None, text=None):
    """References to check, from a manuscript file or from pasted text. Each is a dict with
    n (position), text, doi (or None), title (a guess, or None) and kind ('journal' or 'other')."""
    if text is not None:
        refs = split_reference_text(text)
    else:
        if Path(name or str(path)).suffix.lower() == ".txt":
            refs = split_reference_text(Path(path).read_text(encoding="utf-8", errors="replace"))
        else:
            from step1_split_sections import split_into_sections
            sections, _ = split_into_sections(path, name)
            refs = [t for section, paras in sections if section == "References" for t in paras]
    out = []
    for n, ref in enumerate(refs, start=1):
        doi = find_doi(ref)
        kind = reference_kind(ref, doi)
        out.append({"n": n, "text": ref, "doi": doi, "title": title_guess(ref), "kind": kind,
                    "years": reference_years(ref)})
    return out


def search_query(ref):
    """The text to search Crossref with: the reference without its number, DOI or links."""
    text = strip_number(ref["text"])
    text = DOI_RE.sub(" ", re.sub(r"https?://\S+|doi:\s*", " ", text, flags=I))
    return " ".join(text.split())[:300]


# ---------------------------------------------------------------------------
# Comparing a Crossref record with the reference
# ---------------------------------------------------------------------------
def plain(text):
    text = unicodedata.normalize("NFKD", text or "")
    text = "".join(c for c in text if not unicodedata.combining(c)).lower()
    text = re.sub(r"<[^>]+>", " ", text)                  # Crossref titles can carry <i>…</i>
    return " ".join(re.sub(r"[^a-z0-9]+", " ", text).split())


TITLE_FLAG = re.compile(r"^\s*(?:retracted|withdrawn|retraction)\s*(?:article)?\s*[:\-–]\s*", I)


def record_title(record):
    """The record's title (with subtitle), without a 'RETRACTED:' prefix some publishers add."""
    titles = record.get("title") or []
    title = TITLE_FLAG.sub("", titles[0] if titles else "")
    subtitle = (record.get("subtitle") or [""])[0] if record.get("subtitle") else ""
    return (title + (": " + subtitle if subtitle and subtitle.lower() not in title.lower() else "")).strip()


def record_years(record):
    years = set()
    for key in ("issued", "published-print", "published-online", "published"):
        parts = (record.get(key) or {}).get("date-parts") or [[None]]
        if parts and parts[0] and parts[0][0]:
            years.add(int(parts[0][0]))
    return years


def record_first_author(record):
    for a in record.get("author") or []:
        if a.get("family"):
            return a["family"]
        if a.get("name"):
            return a["name"]
    return ""


def describe(record):
    """'“Title” (Smith, 2020, Journal)' for messages."""
    title = record_title(record)
    bits = [b for b in [record_first_author(record),
                        str(min(record_years(record))) if record_years(record) else "",
                        (record.get("container-title") or [""])[0] if record.get("container-title") else ""] if b]
    short = title if len(title) <= 110 else title[:109].rstrip() + "…"
    return f"“{short}”" + (f" ({', '.join(bits)})" if bits else "")


def title_matches(record, ref):
    """True, False, or None when Crossref has no title to compare."""
    rt = plain(record_title(record))
    if not rt:
        return None
    text = plain(ref["text"])
    main = plain(record_title(record).split(":")[0])
    if rt in text or (len(main.split()) >= 3 and main in text):
        return True
    guess = plain(ref.get("title") or "")
    if guess and (SequenceMatcher(None, guess, rt).ratio() >= 0.8 or SequenceMatcher(None, guess, main).ratio() >= 0.85):
        return True
    words = {w for w in rt.split() if len(w) >= 4}
    if words and len(words & set(text.split())) / len(words) >= 0.8:
        return True
    return False


def author_matches(record, ref):
    surname = plain(record_first_author(record))
    if not surname:
        return None
    text = plain(ref["text"])
    return any(part in text.split() for part in surname.split() if len(part) > 1)


def year_matches(record, ref):
    years = record_years(record)
    if not years or not ref.get("years"):
        return None
    return any(abs(y - r) <= 1 for y in years for r in ref["years"])      # online-first vs print year


# ---------------------------------------------------------------------------
# Retraction and correction notices
# ---------------------------------------------------------------------------
NOTICE_KIND = {"retraction": "retracted", "withdrawal": "retracted", "removal": "retracted", "retracted": "retracted",
               "r": "retracted", "expression_of_concern": "concern", "expression of concern": "concern", "e": "concern",
               "correction": "corrected", "erratum": "corrected", "corrigendum": "corrected", "c": "corrected",
               "reinstatement": "reinstated", "i": "reinstated"}


def date_text(parts):
    try:
        y, m, d = (list(parts[0]) + [None, None])[:3]
        return f"{y}-{m:02d}-{d:02d}" if d else f"{y}-{m:02d}" if m else str(y)
    except Exception:
        return ""


def notices_from_list(entries):
    """Entries from the daily Retraction Watch list: [[code, date, notice DOI], …]."""
    return [{"kind": NOTICE_KIND.get(str(e[0]).lower(), "corrected"), "date": e[1] if len(e) > 1 else "",
             "notice": e[2] if len(e) > 2 else "", "source": "Retraction Watch"} for e in entries or []]


def notices_from_crossref(doi, record=None, update_items=None):
    """Notices about `doi` from Crossref: the record's own 'updated-by' list, and notices found by searching
    for works that update this DOI (filter=updates:<doi>)."""
    found = []
    titles = (record or {}).get("title") or []
    if titles and TITLE_FLAG.match(titles[0]) and not re.match(r"^\s*retraction\b", titles[0], I):
        found.append({"kind": "retracted", "date": "", "notice": "", "source": "the article's own title in Crossref"})
    for u in (record or {}).get("updated-by") or []:
        kind = NOTICE_KIND.get(str(u.get("type", "")).lower())
        if kind:
            found.append({"kind": kind, "date": date_text((u.get("updated") or {}).get("date-parts") or [[]]),
                          "notice": (u.get("DOI") or "").lower(), "source": u.get("source", "Crossref")})
    for item in update_items or []:
        for u in item.get("update-to") or []:
            if (u.get("DOI") or "").lower() == (doi or "").lower():
                kind = NOTICE_KIND.get(str(u.get("type", "")).lower())
                if kind:
                    found.append({"kind": kind, "date": date_text((u.get("updated") or {}).get("date-parts") or [[]]),
                                  "notice": (item.get("DOI") or "").lower(), "source": u.get("source", "Crossref")})
    return found


def strongest_notice(notices):
    """The notice that matters most. A reinstatement after a retraction cancels it."""
    if not notices:
        return None
    by_date = sorted(notices, key=lambda n: n.get("date") or "")
    retractions = [n for n in by_date if n["kind"] == "retracted"]
    reinstated = [n for n in by_date if n["kind"] == "reinstated"]
    if retractions and reinstated and (reinstated[-1].get("date") or "") >= (retractions[-1].get("date") or ""):
        return reinstated[-1]
    for kind in ("retracted", "concern", "corrected", "reinstated"):
        hits = [n for n in by_date if n["kind"] == kind]
        if hits:
            return hits[-1]
    return None


# ---------------------------------------------------------------------------
# The verdict for one reference
# ---------------------------------------------------------------------------
def result(ref, status, details, doi=None, matched=None, notice=None):
    label, severity, fix = STATUSES[status]
    if status == "doi_missing" and doi:
        fix = f"Add the DOI: https://doi.org/{doi}"
    return {"n": ref["n"], "text": ref["text"], "status": status, "label": label, "severity": severity,
            "details": details, "doi": doi or ref.get("doi"), "matched": matched, "notice": notice, "fix": fix}


def best_candidate(ref, items):
    """The search result that is clearly this reference, or None."""
    best, best_score = None, 0
    for item in items or []:
        t, a, y = title_matches(item, ref), author_matches(item, ref), year_matches(item, ref)
        if t and (a or y or (a is None and y is None)):
            score = 2 + bool(a) + bool(y)
            if score > best_score:
                best, best_score = item, score
    return best


def assess(ref, record=None, record_status=None, search_items=None, list_notices=None, update_items=None):
    """The verdict for one reference.
    record / record_status: Crossref's answer for the reference's DOI (status 200 or 404; None if not asked).
    search_items: Crossref's search results, for a reference without a DOI (or whose DOI wasn't found).
    list_notices: entries for the DOI in the daily Retraction Watch list.
    update_items: notices Crossref returned for 'works that update this DOI'."""
    doi = ref.get("doi")
    if doi and record_status not in (200, 404, None):
        return result(ref, "error", f"Crossref didn't answer for {doi} (status {record_status})")

    matched_doi, problem = None, None
    if doi and record_status == 200 and record:
        matched_doi = doi
        t, a, y = title_matches(record, ref), author_matches(record, ref), year_matches(record, ref)
        if t is False and (a is False or y is False):
            problem = ("mismatch", f"The DOI {doi} belongs to {describe(record)}")
        elif t is False:
            problem = ("details", f"The title differs from Crossref's record for this DOI: {describe(record)}")
        elif a is False:
            problem = ("details", f"First author differs: Crossref lists {record_first_author(record)}")
        elif y is False:
            problem = ("details", f"Year differs: Crossref has {min(record_years(record))}")
    elif doi and record_status == 404:
        candidate = best_candidate(ref, search_items)
        if candidate:
            return result(ref, "doi_not_found", f"{doi} isn't a registered DOI; this reference looks like "
                          f"https://doi.org/{candidate['DOI'].lower()}", doi=candidate["DOI"].lower(),
                          matched=describe(candidate))
        return result(ref, "doi_not_found", f"{doi} isn't a registered DOI (check it for typos)")
    elif not doi:
        candidate = best_candidate(ref, search_items)
        if candidate:
            matched_doi, record = candidate["DOI"].lower(), candidate
            problem = ("doi_missing", f"Found in Crossref: {describe(candidate)}")
        elif search_items is None:
            return result(ref, "not_checked", "Not looked up")
        elif ref.get("kind") == "other":
            return result(ref, "not_checked", "Looks like a book, report, guideline or website; not found in Crossref")
        else:
            return result(ref, "not_found", "No matching record in Crossref for this reference")

    notices = notices_from_list(list_notices) + notices_from_crossref(matched_doi, record, update_items)
    notice = strongest_notice(notices)
    if notice:
        what = {"retracted": "Retracted", "concern": "Expression of concern", "corrected": "Corrected",
                "reinstated": "Reinstated after retraction"}[notice["kind"]]
        when = f" on {notice['date']}" if notice.get("date") else ""
        source = {"publisher": "the publisher (via Crossref)", "retraction-watch": "Retraction Watch (via Crossref)",
                  "Crossref": "Crossref"}.get(notice["source"], notice["source"])
        link = f" (notice: https://doi.org/{notice['notice']})" if notice.get("notice") else ""
        extra = f". Also: {problem[1]}" if problem and problem[0] != "details" else ""
        return result(ref, notice["kind"], f"{what}{when}{link}, according to {source}{extra}",
                      doi=matched_doi, matched=describe(record) if record else None, notice=notice)
    if problem:
        return result(ref, problem[0], problem[1], doi=matched_doi, matched=describe(record) if record else None)
    return result(ref, "verified", f"Matches Crossref: {describe(record)}", doi=matched_doi, matched=describe(record))


def summary(results):
    counts = {}
    for r in results:
        counts[r["status"]] = counts.get(r["status"], 0) + 1
    problems = sum(1 for r in results if r["severity"] in ("must", "should"))
    if problems:
        return counts, "Needs revision"
    return counts, ("Incomplete: some references couldn't be checked" if counts.get("error") else "Ready")


# ---------------------------------------------------------------------------
# Reports
# ---------------------------------------------------------------------------
FILL = {"must": "F8CBAD", "should": "FFE699", "advisory": "DDEBF7", "info": "EDEDED", "ok": "C6EFCE"}


def build_excel(results, source_name=""):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "References"
    counts, verdict = summary(results)
    ws["A1"], ws["B1"] = "File", source_name or "pasted references"
    ws["A2"], ws["B2"] = "Result", verdict
    ws["A3"], ws["B3"] = "References", f"{len(results)} checked: " + ", ".join(
        f"{n} {STATUSES[s][0].lower()}" for s, n in sorted(counts.items(), key=lambda kv: SEVERITY_ORDER[STATUSES[kv[0]][1]]))
    for r in range(1, 4):
        ws.cell(r, 1).font = Font(bold=True)
    headers = ["#", "Reference", "Result", "Severity", "Details", "DOI", "How to fix"]
    for col, h in enumerate(headers, start=1):
        c = ws.cell(5, col, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
    for row, r in enumerate(sorted(results, key=lambda r: (SEVERITY_ORDER[r["severity"]], r["n"])), start=6):
        values = [r["n"], r["text"], r["label"], SEVERITY_LABEL[r["severity"]], r["details"],
                  f"https://doi.org/{r['doi']}" if r.get("doi") else "", r["fix"]]
        for col, v in enumerate(values, start=1):
            c = ws.cell(row, col, v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row, 3).fill = PatternFill("solid", fgColor=FILL[r["severity"]])
    for col, width in zip("ABCDEFG", [5, 60, 22, 12, 60, 34, 50]):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A6"
    return wb


def build_word(results, source_name=""):
    from datetime import date
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Inches, Pt, RGBColor
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading("Reference integrity report", 0)
    counts, verdict = summary(results)
    doc.add_paragraph(f"{source_name or 'Pasted references'}    ·    {date.today():%d %B %Y}    ·    Result: {verdict}. "
                      "Checked against Crossref and Retraction Watch; a screening aid for a person to review.")
    colours = {"must": RGBColor(0xC0, 0x39, 0x2B), "should": RGBColor(0xB7, 0x79, 0x0F),
               "advisory": RGBColor(0x2E, 0x6D, 0xA4), "info": RGBColor(0x70, 0x70, 0x70),
               "ok": RGBColor(0x2E, 0x7D, 0x32)}
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    for cell, text in zip(table.rows[0].cells, ["#", "Reference", "Result", "What to do"]):
        cell.text = text
    for r in sorted(results, key=lambda r: (SEVERITY_ORDER[r["severity"]], r["n"])):
        row = table.add_row().cells
        row[0].text = str(r["n"])
        row[1].text = r["text"]
        run = row[2].paragraphs[0].add_run(r["label"])
        run.bold = True
        run.font.color.rgb = colours[r["severity"]]
        row[2].add_paragraph(r["details"])
        row[3].text = r["fix"]
    for row in table.rows:
        for cell, width in zip(row.cells, (Inches(0.4), Inches(3.8), Inches(2.8), Inches(2.5))):
            cell.width = width
    return doc


# ---------------------------------------------------------------------------
# Command line: live lookups with urllib
# ---------------------------------------------------------------------------
class Crossref:
    """Polite, rate-limited Crossref lookups (at most ~8 single-record and ~2 search requests a second)."""

    def __init__(self, contact=CONTACT):
        self.contact = contact
        self.last = 0.0
        self.offline = False          # set when Crossref can't be reached at all: stop trying

    def get(self, url, gap):
        if self.offline:
            return 0, None
        wait = self.last + gap - time.time()
        if wait > 0:
            time.sleep(wait)
        request = urllib.request.Request(url, headers={
            "User-Agent": f"GreySwanReferenceChecker/1.0 (https://thegreyswan.in/references/; mailto:{self.contact})"})
        for attempt in range(4):
            self.last = time.time()
            try:
                with urllib.request.urlopen(request, timeout=30) as response:
                    return 200, json.loads(response.read().decode("utf-8"))
            except urllib.error.HTTPError as e:
                if e.code == 429 or e.code >= 500:
                    time.sleep(2 * (attempt + 1))
                    continue
                return e.code, None
            except (urllib.error.URLError, TimeoutError, OSError) as e:
                if attempt >= 1:
                    self.offline = True
                    print(f"  Can't reach Crossref ({getattr(e, 'reason', e)}). Check the internet connection.")
                    return 0, None
                time.sleep(2)
        return 0, None

    def work(self, doi):
        status, data = self.get(f"{API}/{urllib.parse.quote(doi, safe='/:;()')}?mailto={self.contact}", 0.13)
        return status, (data or {}).get("message")

    def search(self, query):
        url = (f"{API}?query.bibliographic={urllib.parse.quote(query)}&rows=3&select={SEARCH_FIELDS}"
               f"&mailto={self.contact}")
        status, data = self.get(url, 0.4)
        return ((data or {}).get("message") or {}).get("items") if status == 200 else None

    def updates(self, doi):
        status, data = self.get(f"{API}?filter=updates:{urllib.parse.quote(doi, safe='/:;()')}&rows=10"
                                f"&mailto={self.contact}", 0.4)
        return ((data or {}).get("message") or {}).get("items") if status == 200 else None


def load_retraction_list(path=None):
    """The daily Retraction Watch list ({doi: [[code, date, notice DOI], …]}), or None if unavailable."""
    try:
        if path:
            raw = Path(path).read_bytes()
        else:
            with urllib.request.urlopen(RETRACTIONS_URL, timeout=60) as response:
                raw = response.read()
        return json.loads(gzip.decompress(raw).decode("utf-8"))["notices"]
    except Exception as e:
        print(f"  (Retraction Watch list unavailable: {e}; using Crossref's notices instead)")
        return None


def check_all(refs, crossref, retraction_list, progress=print):
    results = []
    for ref in refs:
        record, status, items, updates = None, None, None, None
        if getattr(crossref, "offline", False):
            results.append(result(ref, "error", "Not checked: Crossref couldn't be reached"))
            continue
        if ref["doi"]:
            status, record = crossref.work(ref["doi"])
            if status == 404:
                items = crossref.search(search_query(ref))
        else:
            items = crossref.search(search_query(ref))
            if items is None:                   # the search itself failed: say so, don't call it "not found"
                results.append(result(ref, "error", "Crossref didn't answer the search"))
                progress(f"  {ref['n']:>3}. Couldn't check")
                continue
        doi = ref["doi"] if status == 200 else None
        if not doi and items:
            candidate = best_candidate(ref, items)
            doi = candidate["DOI"].lower() if candidate else None
        if doi and retraction_list is None:
            updates = crossref.updates(doi)
        found = assess(ref, record, status, items, (retraction_list or {}).get(doi or "", []), updates)
        results.append(found)
        progress(f"  {ref['n']:>3}. {found['label']}: {found['details'][:100]}")
    return results


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if not args:
        print(__doc__.strip().split("How to run")[1])
        return
    path = Path(args[0])
    refs = extract(path, path.name)
    if not refs:
        print("No references found. Is there a 'References' heading?")
        return
    print(f"{len(refs)} references in {path.name}. Checking against Crossref…")
    retraction_list = load_retraction_list(sys.argv[sys.argv.index("--retractions") + 1]
                                           if "--retractions" in sys.argv else None)
    results = check_all(refs, Crossref(), retraction_list)
    counts, verdict = summary(results)
    print(f"\n{verdict}: " + ", ".join(f"{n} {STATUSES[s][0].lower()}" for s, n in counts.items()))
    build_excel(results, path.name).save("reference_report.xlsx")
    print("Report saved: reference_report.xlsx")
    if "--word" in sys.argv:
        build_word(results, path.name).save("reference_report.docx")
        print("Word report saved: reference_report.docx")


if __name__ == "__main__":
    main()
