"""
Builds the labelled test set in tests/papers/.

Each paper is invented and has a known list of planted problems, written to
tests/expected.json as "Section|Check" labels. run_tests.py compares the
checker's findings against these labels.

    python tests/make_test_papers.py
"""

import json
from pathlib import Path
from docx import Document

OUT = Path(__file__).parent / "papers"
OUT.mkdir(exist_ok=True)
EXPECTED = {}

STANDARD_BACK_MATTER = [
    ("Acknowledgements", ["We thank the participating staff for their support."]),
    ("Funding", ["This study received no specific funding."]),
    ("Conflicts of interest", ["The authors declare no conflicts of interest."]),
    ("Author contributions", ["AB: conceptualisation, data collection. CD: analysis, writing."]),
    ("Data availability statement", ["Data are available from the corresponding author on reasonable request."]),
]


def numbered_refs(n):
    return [f"{i}. Author{i} A, Writer B. Example study {i}. J Example Med. 202{i % 5};{i}:1-9." for i in range(1, n + 1)]


def build(name, title, sections, expected, tables=0):
    """sections: list of (heading, [paragraphs]); heading None = plain paragraphs."""
    d = Document()
    d.add_paragraph(title)
    d.add_paragraph("A. Bose, C. Das. Department of Community Medicine, Example Medical College, Ghaziabad, India.")
    for heading, paras in sections:
        if heading:
            d.add_heading(heading, 1)
        for p in paras:
            d.add_paragraph(p)
    for _ in range(tables):
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "Variable", "Value"
    d.save(OUT / name)
    EXPECTED[name] = sorted(expected)


# ---------------------------------------------------------------------------
# 1. CONTROL: a complete, well-reported randomised trial. Should come back clean.
build("01_control_clean_rct.docx",
      "Effect of a structured walking programme on glycaemic control in adults with type 2 diabetes: a randomised controlled trial",
      [("Abstract", ["Background: Physical activity improves glycaemic control. Methods: In this randomised controlled trial, "
                     "184 adults with type 2 diabetes were allocated to a walking programme or usual care for 24 weeks. "
                     "Results: Glycated haemoglobin fell by 0.6% more with walking. Conclusion: Structured walking is effective."]),
       ("Keywords", ["type 2 diabetes; exercise; randomised controlled trial; glycaemic control"]),
       ("Introduction", ["Type 2 diabetes affects millions of adults in India [1]. Physical activity lowers blood glucose [2,3]."]),
       ("Methods", ["This parallel-group randomised controlled trial was conducted between January 2024 and December 2024 at the "
                    "diabetes clinic of Example Medical College and is reported according to CONSORT. The trial was registered with "
                    "the Clinical Trials Registry of India (CTRI/2023/12/060123).",
                    "A total of 184 participants (n = 184) were enrolled. Inclusion criteria were age 30–65 years and a diagnosis of "
                    "type 2 diabetes for at least one year; exclusion criteria were insulin therapy and inability to walk unaided.",
                    "The study was approved by the Institutional Ethics Committee of Example Medical College (approval no. IEC/2023/211, "
                    "dated 14 November 2023). Written informed consent was obtained from all participants.",
                    "Body mass index (BMI) and glycated haemoglobin were measured at baseline and 24 weeks [4]."]),
       ("Results", ["Baseline characteristics were similar between groups (Table 1). Glycated haemoglobin fell more with walking "
                    "(Figure 1), and adherence was high (Figure 2; Table 2)."]),
       ("Discussion", ["Our findings agree with earlier trials of supervised exercise [5]. BMI changes were modest [6]."]),
       *STANDARD_BACK_MATTER,
       (None, ["Figure 1. Change in glycated haemoglobin by group.", "Figure 2. Weekly adherence.",
               "Table 1. Baseline characteristics.", "Table 2. Adherence by month."]),
       ("References", numbered_refs(6))],
      expected=[], tables=2)

# ---------------------------------------------------------------------------
# 2. Flawed RCT: six planted problems.
build("02_flawed_rct.docx",
      "Vitamin D supplementation and kidney function in chronic kidney disease: a randomised controlled trial",
      [("Abstract", ["Background: Vitamin D deficiency is common in chronic kidney disease. Methods: In this randomised controlled trial, "
                     "120 patients received vitamin D or placebo for 12 weeks. Results: Kidney function was unchanged. "
                     "Conclusion: Supplementation did not improve kidney function."]),
       ("Keywords", ["chronic kidney disease; vitamin D; randomised controlled trial"]),
       ("Introduction", ["Chronic kidney disease is a growing public health problem [1,2]."]),
       ("Methods", ["This randomised controlled trial was conducted between March 2024 and February 2025 at Example Hospital. "
                    "A total of 120 patients were enrolled. Inclusion criteria were stage 3 disease; exclusion criteria were dialysis "
                    "and pregnancy.",                                                        # PLANTED: no trial registration
                    "The study was approved by the Institutional Ethics Committee of Example Hospital (approval no. IEC/2024/045).",  # PLANTED: no approval date
                    "Verbal consent was obtained from all patients.",                       # PLANTED: verbal consent only
                    "The eGFR was calculated at baseline and 12 weeks [3]."]),              # PLANTED: eGFR never defined
       ("Results", ["Kidney function over time is shown in Figure 2 and Figure 1 [9]."]),   # PLANTED: figure order; ref 9 not in list
       ("Discussion", ["Similar null results were reported earlier [4,5]."]),
       *STANDARD_BACK_MATTER,
       (None, ["Figure 1. Study flow.", "Figure 2. Kidney function over time."]),
       ("References", numbered_refs(6))],
      expected=["Study details|Trial registration", "Declarations|Ethics approval date", "Declarations|Verbal consent",
                "Abbreviations (Main text)|Not defined", "Figures & tables|Figure order", "References|Citation not in list",
                "References|Reference not cited"])                                      # reference 6 is genuinely never cited

# ---------------------------------------------------------------------------
# 3. Systematic review: no PROSPERO number, abstract too long, author-year references.
long_abstract = ("Background: Antimicrobial stewardship programmes aim to improve antibiotic use in hospitals. " * 10 +
                 "Methods: We searched four databases for studies published between 2010 and 2025 and followed the PRISMA guideline. " * 5 +
                 "Results: Twenty-eight studies met the inclusion criteria and showed reduced antibiotic consumption. " * 5 +
                 "Conclusion: Stewardship programmes reduce antibiotic use. " * 4)          # about 280 words, over the 250 limit
build("03_systematic_review.docx",
      "Antimicrobial stewardship programmes in Indian hospitals: a systematic review and meta-analysis",
      [("Abstract", [long_abstract]),                                                       # PLANTED: abstract over 250 words
       ("Keywords", ["antimicrobial stewardship; systematic review; antibiotic use"]),
       ("Introduction", ["Antibiotic resistance is a major threat (Laxminarayan et al., 2016)."]),
       ("Methods", ["We conducted a systematic review and meta-analysis following PRISMA 2020. MEDLINE, Embase, Scopus and "
                    "IndMED were searched for studies published between January 2010 and June 2025. Two reviewers screened "
                    "records independently."]),                                             # PLANTED: no PROSPERO number
       ("Results", ["Twenty-eight studies were included (Figure 1). Pooled results are shown in Table 1."]),
       ("Discussion", ["Our findings agree with earlier reviews (Davey et al., 2017)."]),
       *STANDARD_BACK_MATTER,
       (None, ["Figure 1. PRISMA flow diagram.", "Table 1. Characteristics of included studies."]),
       ("References", ["Davey P, Marwick CA, Scott CL. Interventions to improve antibiotic prescribing. Cochrane Database Syst Rev. 2017.",
                       "Laxminarayan R, Matsoso P, Pant S. Access to effective antimicrobials. Lancet. 2016;387:168-75."])],
      expected=["Abstract|Abstract length", "Study details|PROSPERO registration"], tables=1)

# ---------------------------------------------------------------------------
# 4. Animal study: no euthanasia statement, data from 2018–2019.
build("04_animal_study.docx",
      "Hepatoprotective effect of a plant extract in a rat model of paracetamol-induced liver injury",
      [("Abstract", ["Background: Paracetamol overdose causes liver injury. Methods: Thirty Wistar rats received the extract or vehicle. "
                     "Results: Liver enzymes were lower in treated rats. Conclusion: The extract protects the liver."]),
       ("Keywords", ["hepatoprotection; paracetamol; Wistar rats"]),
       ("Introduction", ["Drug-induced liver injury is common [1]. Plant extracts may protect the liver [2]."]),
       ("Methods", ["Thirty male Wistar rats (n = 30) were studied between June 2018 and March 2019.",  # PLANTED: data before 2020
                    "The protocol was approved by the Institutional Animal Ethics Committee (IAEC) of Example College "
                    "(approval no. IAEC/2018/07, dated 2 May 2018) and followed CPCSEA guidelines.",
                    "Alanine aminotransferase (ALT) was measured on day 7 [3]. ALT values were compared by one-way ANOVA."]),
       ("Results", ["ALT was lower in treated rats (Figure 1; Table 1)."]),          # PLANTED: no euthanasia statement anywhere
       ("Discussion", ["These results agree with earlier work [4]."]),
       *STANDARD_BACK_MATTER,
       (None, ["Figure 1. Serum ALT by group.", "Table 1. Body weights."]),
       ("References", numbered_refs(4))],
      expected=["Study details|Euthanasia", "Study details|Data older than 2020"], tables=1)

# ---------------------------------------------------------------------------
# 5. School survey: minors without parental consent, questionnaire not supplied, no inclusion/exclusion.
build("05_survey_minors.docx",
      "Screen time and sleep among school students in Ghaziabad: a cross-sectional survey",
      [("Abstract", ["Background: Screen time may disturb sleep. Methods: A total of 412 school students aged 12–16 years completed a "
                     "questionnaire. Results: Longer screen time was linked to shorter sleep. Conclusion: Screen time guidance is needed."]),
       ("Keywords", ["screen time; sleep; adolescents; survey"]),
       ("Introduction", ["Adolescents increasingly use smartphones [1,2]."]),
       ("Methods", ["This cross-sectional survey was conducted between July 2024 and September 2024 in four schools. "
                    "A total of 412 school students aged 12–16 years participated.",              # PLANTED: no inclusion/exclusion
                    "Students completed a self-administered questionnaire on screen time and sleep.",   # PLANTED: questionnaire not supplied
                    "The study was approved by the Institutional Ethics Committee of Example Medical College "
                    "(approval no. IEC/2024/088, dated 3 June 2024). Written informed consent was obtained from all students."]),  # PLANTED: no parental consent
       ("Results", ["Mean screen time was 4.2 hours (Table 1)."]),
       ("Discussion", ["Similar patterns were reported elsewhere [3]."]),
       *STANDARD_BACK_MATTER,
       (None, ["Table 1. Screen time by age group."]),
       ("References", numbered_refs(3))],
      expected=["Declarations|Consent for minors", "Study details|Questionnaire", "Study details|Inclusion/exclusion criteria"], tables=1)

# ---------------------------------------------------------------------------
# 6. Case report: no IMRaD by design. A good checker should not ask for Methods or Results.
build("06_case_report.docx",
      "An unusual presentation of scrub typhus with acute pancreatitis: a case report",
      [("Abstract", ["Scrub typhus rarely causes pancreatitis. We describe a 34-year-old man with fever and abdominal pain who "
                     "was diagnosed with scrub typhus and acute pancreatitis and recovered with doxycycline."]),
       ("Keywords", ["scrub typhus; pancreatitis; case report"]),
       ("Introduction", ["Scrub typhus is endemic in India [1]."]),
       ("Case presentation", ["A 34-year-old man presented in March 2025 with fever for six days and epigastric pain. "
                              "Serum lipase was raised and an eschar was found on the abdomen (Figure 1). "
                              "Scrub typhus IgM was positive [2]. He recovered with doxycycline."]),
       ("Discussion", ["Pancreatitis is a rare complication of scrub typhus [3]."]),
       ("Conclusion", ["Clinicians should consider scrub typhus in febrile patients with abdominal pain."]),
       ("Consent", ["Written informed consent for publication was obtained from the patient."]),
       ("Ethics", ["Ethical approval is not required for a single case report at our institution."]),
       *STANDARD_BACK_MATTER,
       (None, ["Figure 1. Eschar on the abdominal wall."]),
       ("References", numbered_refs(3))],
      expected=[])

# ---------------------------------------------------------------------------
# 7. Messy but correct headings: a clean paper written in a different journal style.
build("07_messy_headings.docx",
      "Prevalence of anaemia among pregnant women attending antenatal clinics: a cross-sectional study",
      [("ABSTRACT", ["Background: Anaemia in pregnancy remains common. Methods: We studied 350 pregnant women. "
                     "Results: Anaemia was found in 46%. Conclusion: Screening should continue."]),
       ("Key words:", ["anaemia; pregnancy; antenatal care"]),
       ("1. BACKGROUND", ["Anaemia affects many pregnant women in India [1]."]),
       ("2. MATERIAL AND METHODS", ["This cross-sectional study was conducted between February 2024 and July 2024 at the antenatal "
                                    "clinic of Example Medical College. A total of 350 pregnant women were enrolled. Inclusion "
                                    "criteria were singleton pregnancy; exclusion criteria were known haemoglobinopathy.",
                                    "The Institutional Ethics Committee approved the study (Ref. No. IEC/2024/019, dated 10 January 2024). "
                                    "Written informed consent was obtained from all participants."]),
       ("2.1 Statistical analysis", ["Proportions were compared with the chi-square test."]),
       ("3. RESULTS AND DISCUSSION", ["Anaemia was found in 46% of women (Table 1). This matches national data [2]."]),
       ("Declaration of competing interest", ["The authors declare that they have no competing interests."]),
       ("Funding sources", ["None."]),
       ("CRediT authorship contribution statement", ["AB: Conceptualization, Investigation. CD: Formal analysis, Writing."]),
       ("Acknowledgments", ["We thank the clinic nurses."]),
       ("Data availability", ["Data are available on reasonable request."]),
       (None, ["Table 1. Anaemia by trimester."]),
       ("REFERENCES", numbered_refs(2))],
      expected=[], tables=1)

(Path(__file__).parent / "expected.json").write_text(json.dumps(EXPECTED, indent=2), encoding="utf-8")
print(f"Built {len(EXPECTED)} test papers in {OUT}")


# ===========================================================================
# ROUND 2: unseen papers, labels written before the checker was run on them.
# ===========================================================================
from docx.shared import Pt  # noqa: E402

# 8. Retrospective chart review: valid consent waiver by the committee; data from 2017 (planted).
build("08_retrospective_waiver.docx",
      "Outcomes of dengue with warning signs in a tertiary hospital: a retrospective study",
      [("Abstract", ["Background: Dengue with warning signs can progress to severe disease. Methods: We reviewed the records of "
                     "268 patients. Results: Thrombocytopenia predicted ICU admission. Conclusion: Early monitoring is essential."]),
       ("Keywords", ["dengue; warning signs; retrospective study"]),
       ("Introduction", ["Dengue is a major cause of hospital admission in India [1]."]),
       ("Methods", ["We conducted a retrospective review of medical records of 268 patients admitted between January 2017 and "
                    "December 2017 at Example Hospital.",                                    # PLANTED: data before 2020
                    "Inclusion criteria were laboratory-confirmed dengue; exclusion criteria were incomplete records.",
                    "The Institutional Ethics Committee of Example Hospital approved the study (approval no. IEC/2023/301, "
                    "dated 9 August 2023) and waived the requirement for informed consent because of the retrospective design.",
                    "Intensive care unit (ICU) admission was the primary outcome [2]."]),
       ("Results", ["ICU admission occurred in 31 patients (Table 1)."]),
       ("Discussion", ["Earlier studies reported similar ICU rates [3]."]),
       *STANDARD_BACK_MATTER,
       (None, ["Table 1. Predictors of ICU admission."]),
       ("References", numbered_refs(3))],
      expected=["Study details|Data older than 2020",
                "Abbreviations (Abstract)|Not defined"], tables=1)              # ICU is used in the abstract without definition

# 9. Qualitative interview study: interview guide supplied as a supplement; written consent. Clean.
build("09_qualitative_interviews.docx",
      "Experiences of community health workers during the COVID-19 pandemic: a qualitative study",
      [("Abstract", ["Background: Community health workers carried a heavy load during the pandemic. Methods: We conducted "
                     "in-depth interviews with 24 workers. Results: Three themes emerged. Conclusion: Workers need better support."]),
       ("Keywords", ["community health workers; qualitative; COVID-19"]),
       ("Introduction", ["Community health workers are central to primary care in India [1]."]),
       ("Methods", ["We conducted semi-structured interviews with 24 community health workers between April 2024 and June 2024. "
                    "Inclusion criteria were at least two years of service; exclusion criteria were absence during the pandemic.",
                    "The interview guide is provided in Supplementary File 1. The study follows COREQ.",
                    "The Institutional Ethics Committee of Example Medical College approved the study (approval no. IEC/2024/066, "
                    "dated 18 March 2024). Written informed consent was obtained from all participants."]),
       ("Results", ["Three themes emerged (Table 1)."]),
       ("Discussion", ["Our findings echo earlier reports [2]."]),
       *STANDARD_BACK_MATTER,
       (None, ["Table 1. Themes and illustrative quotes."]),
       ("References", numbered_refs(2))],
      expected=[], tables=1)

# 10. Narrative review: no Methods or Results by design. Clean.
build("10_narrative_review.docx",
      "Antimicrobial resistance in Indian agriculture: a narrative review",
      [("Abstract", ["Antibiotic use in agriculture contributes to resistance. This narrative review summarises evidence on "
                     "antibiotic use in Indian poultry and aquaculture and outlines policy options."]),
       ("Keywords", ["antimicrobial resistance; agriculture; narrative review"]),
       ("Introduction", ["Antibiotics are widely used in food animals [1]."]),
       ("Antibiotic use in poultry", ["Colistin use in poultry has been reported widely [2]."]),
       ("Aquaculture", ["Residues have been detected in farmed fish [3]."]),
       ("Policy options", ["India banned colistin as a growth promoter in 2019 [4]."]),
       ("Conclusion", ["Stronger surveillance is needed."]),
       *STANDARD_BACK_MATTER,
       ("References", numbered_refs(4))],
      expected=[])

# 11. Letter to the editor: short, no abstract or IMRaD. Clean.
build("11_letter_to_editor.docx",
      "Letter to the Editor: Rising scrub typhus cases in northern India",
      [(None, ["Dear Editor,",
               "We read with interest the recent article on scrub typhus [1]. Our hospital has seen a rise in cases in 2025, "
               "which suggests the disease is spreading beyond its usual areas [2].",
               "We suggest wider testing in febrile patients."]),
       *STANDARD_BACK_MATTER,
       ("References", numbered_refs(2))],
      expected=[])

# 12. Clean cohort study whose citations are superscript numbers, not [1] brackets. Clean.
d = Document()
d.add_paragraph("Tuberculosis treatment outcomes among people with diabetes: a prospective cohort study")
d.add_paragraph("A. Bose, C. Das. Department of Pulmonary Medicine, Example Medical College, Ghaziabad, India.")
def sup_para(text_parts):
    p = d.add_paragraph()
    for part, sup in text_parts:
        r = p.add_run(part); r.font.superscript = sup
for heading, parts in [
    ("Abstract", [[("Background: Diabetes may worsen tuberculosis outcomes. Methods: We followed 220 patients for six months. "
                    "Results: Treatment failure was higher with diabetes. Conclusion: Glycaemic control matters.", False)]]),
    ("Keywords", [[("tuberculosis; diabetes; cohort study", False)]]),
    ("Introduction", [[("Diabetes triples the risk of tuberculosis.", False), ("1,2", True)]]),
    ("Methods", [[("This prospective cohort study enrolled 220 patients between May 2024 and April 2025 at Example Medical College. "
                   "Inclusion criteria were new pulmonary tuberculosis; exclusion criteria were drug resistance.", False)],
                 [("The Institutional Ethics Committee approved the study (approval no. IEC/2024/030, dated 2 April 2024). "
                   "Written informed consent was obtained from all participants.", False)]]),
    ("Results", [[("Treatment failure was more common with diabetes (Table 1).", False), ("3", True)]]),
    ("Discussion", [[("This agrees with earlier cohorts.", False), ("4", True)]]),
]:
    d.add_heading(heading, 1)
    for parts_ in parts:
        sup_para(parts_)
for h, t in STANDARD_BACK_MATTER:
    d.add_heading(h, 1); d.add_paragraph(t[0])
d.add_paragraph("Table 1. Outcomes by diabetes status.")
t = d.add_table(rows=2, cols=2); t.cell(0, 0).text = "Outcome"
d.add_heading("References", 1)
for r in numbered_refs(4):
    d.add_paragraph(r)
d.save(OUT / "12_superscript_citations.docx")
EXPECTED["12_superscript_citations.docx"] = []

(Path(__file__).parent / "expected.json").write_text(json.dumps(EXPECTED, indent=2), encoding="utf-8")
print(f"Round 2 added; {len(EXPECTED)} papers in total")
