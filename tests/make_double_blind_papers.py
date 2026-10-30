"""
Builds the double-blind test papers (13 and 14) and their expected results.

  13_double_blind_clean.docx  properly blinded: should come back clean under every profile
  14_double_blind_leaky.docx  gives the authors away in eight different ways

Also writes tests/expected_double_blind.json: the EXTRA problems each paper should show when
checked against a double-blind profile (most older test papers include a title page and name
their institution, which is fine in single-blind review but not in double-blind).

    python tests/make_double_blind_papers.py      (after make_test_papers.py)
"""

import json
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).parent
OUT = HERE / "papers"


def refs(n):
    return [f"{i}. Author{i} A, Writer B. Example study {i}. J Example Med. 202{i % 5};{i}:1-9." for i in range(1, n + 1)]


def add_sections(d, sections):
    for heading, paras in sections:
        if heading:
            d.add_heading(heading, 1)
        for p in paras:
            d.add_paragraph(p)


ABSTRACT = [
    ("Abstract", ["Background: Hand hygiene reduces hospital infections. Methods: In this randomised controlled trial, 160 nurses "
                  "were allocated to feedback or usual training for 12 weeks. Results: Compliance rose by 18% with feedback. "
                  "Conclusion: Feedback improves hand hygiene."]),
    ("Keywords", ["hand hygiene; nurses; infection control; randomised controlled trial"]),
    ("Introduction", ["Hospital infections are common in India [1]. Hand hygiene is the simplest prevention [2,3]."]),
]
RESULTS = [
    ("Results", ["Groups were similar at baseline (Table 1). Compliance rose more with feedback (Figure 1)."]),
]
FIGURES = [(None, ["Figure 1. Hand hygiene compliance by group.", "Table 1. Baseline characteristics."])]

# ---------------------------------------------------------------------------
# 13. Properly blinded: title only, placeholders, third-person self-citation, clean file properties
d = Document()
d.core_properties.author = ""
d.core_properties.last_modified_by = ""
d.add_paragraph("Audit and feedback to improve hand hygiene among nurses: a randomised controlled trial")
add_sections(d, ABSTRACT + [
    ("Methods", ["This randomised controlled trial was conducted between January 2024 and June 2024 in a tertiary care hospital "
                 "in northern India and is reported according to CONSORT. It was registered with the Clinical Trials Registry of "
                 "India (CTRI/2023/11/059876).",
                 "A total of 160 nurses (n = 160) took part. Inclusion criteria were permanent ward staff; exclusion criteria were "
                 "planned leave during the study.",
                 "The study was approved by the Institutional Ethics Committee of [blinded for review] (approval no. IEC/2023/187, "
                 "dated 6 December 2023). Written informed consent was obtained from all participants."]),
    *RESULTS,
    ("Discussion", ["Bose et al. previously reported similar gains in intensive care [4]. Feedback was well accepted [5]."]),
    ("Acknowledgements", ["[Removed for peer review]"]),
    ("Funding", ["This study received no specific funding."]),
    ("Conflicts of interest", ["The authors declare no conflicts of interest."]),
    ("Author contributions", ["[Blinded for review]"]),
    ("Data availability statement", ["Data are available from the corresponding author on reasonable request."]),
    *FIGURES,
    ("References", refs(5)),
])
d.add_table(rows=2, cols=2).cell(0, 0).text = "Variable"
d.save(OUT / "13_double_blind_clean.docx")

# ---------------------------------------------------------------------------
# 14. Leaky: title page, header, file properties, tracked change, comment, email in text,
#     self-citation, acknowledgements naming a person, institution in Methods, names in contributions
d = Document()
d.core_properties.author = "Priya Sharma"
d.core_properties.last_modified_by = "Rahul Verma"
d.sections[0].header.paragraphs[0].text = "Priya Sharma et al. | Lakeview Medical College"
d.add_paragraph("Audit and feedback to improve hand hygiene among nurses: a randomised controlled trial")
d.add_paragraph("Priya Sharma1*, Rahul Verma2")
d.add_paragraph("1 Department of Microbiology, Lakeview Medical College, Pune, India. 2 School of Nursing, Lakeview University, Pune, India.")
d.add_paragraph("*Corresponding author: priya.sharma@lakeview.edu.in; ORCID 0000-0002-1825-0097")
add_sections(d, ABSTRACT + [
    ("Methods", ["This randomised controlled trial was conducted between January 2024 and June 2024 at Lakeview Medical College "
                 "and is reported according to CONSORT. It was registered with the Clinical Trials Registry of India "
                 "(CTRI/2023/11/059876).",
                 "A total of 160 nurses (n = 160) took part. Inclusion criteria were permanent ward staff; exclusion criteria were "
                 "planned leave during the study.",
                 "The study was approved by the Institutional Ethics Committee (approval no. IEC/2023/187, dated 6 December 2023). "
                 "Written informed consent was obtained from all participants."]),
    *RESULTS,
    ("Discussion", ["In our previous study, we found similar gains in intensive care [4]. Feedback was well accepted [5]."]),
    ("Acknowledgements", ["We thank Dr. Anil Mehta for help with data collection."]),
    ("Funding", ["This study received no specific funding."]),
    ("Conflicts of interest", ["The authors declare no conflicts of interest."]),
    ("Author contributions", ["Priya Sharma: conceptualisation, writing. Rahul Verma: data collection, analysis."]),
    ("Data availability statement", ["Data are available from the corresponding author (priya.sharma@lakeview.edu.in) "
                                     "on reasonable request."]),
    *FIGURES,
    ("References", refs(5)),
])
d.add_table(rows=2, cols=2).cell(0, 0).text = "Variable"

# a tracked insertion by the second author, and a comment by the first
results_para = next(p for p in d.paragraphs if p.text.startswith("Groups were similar"))
ins = OxmlElement("w:ins")
ins.set(qn("w:id"), "101"); ins.set(qn("w:author"), "Rahul Verma"); ins.set(qn("w:date"), "2026-09-01T10:00:00Z")
run = OxmlElement("w:r"); text = OxmlElement("w:t"); text.text = " Adherence was high."
text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
run.append(text); ins.append(run); results_para._p.append(ins)
d.add_comment(results_para.runs[0], text="Check these numbers again", author="Priya Sharma", initials="PS")
d.save(OUT / "14_double_blind_leaky.docx")

# ---------------------------------------------------------------------------
# Expected results
expected_path = HERE / "expected.json"
expected = json.loads(expected_path.read_text(encoding="utf-8"))
expected["13_double_blind_clean.docx"] = []
expected["14_double_blind_leaky.docx"] = ["Structure|Tracked changes / comments"]
expected_path.write_text(json.dumps(expected, indent=2), encoding="utf-8")

TITLE_PAGE = "Anonymisation|Title page in manuscript"
INSTITUTION = "Anonymisation|Institution named"
extra = {name: [TITLE_PAGE] for name in expected}           # papers 1-12 all include a title page
for name in ["01_control_clean_rct.docx", "02_flawed_rct.docx", "04_animal_study.docx", "05_survey_minors.docx",
             "07_messy_headings.docx", "08_retrospective_waiver.docx", "09_qualitative_interviews.docx",
             "12_superscript_citations.docx"]:
    extra[name].append(INSTITUTION)                         # their Methods name the study site
extra["13_double_blind_clean.docx"] = []
extra["14_double_blind_leaky.docx"] = [
    TITLE_PAGE, INSTITUTION,
    "Anonymisation|Emails / ORCID iDs",
    "Anonymisation|Author names in text",
    "Anonymisation|Headers / footers",
    "Anonymisation|File properties",
    "Anonymisation|Self-citation",
    "Anonymisation|Acknowledgements",
]
(HERE / "expected_double_blind.json").write_text(
    json.dumps({k: sorted(v) for k, v in extra.items()}, indent=2), encoding="utf-8")
print("Built papers 13 and 14; wrote expected.json and expected_double_blind.json")
