"""Builds a realistic test manuscript full of things that caused false alarms."""
from docx import Document
d = Document()
d.add_paragraph("EXPRESSION OF TP53 AND BRCA1 IN BREAST CANCER CELL LINES TREATED WITH A NOVEL COMPOUND")
d.add_paragraph("R. Sharma, K. Iyer. Department of Biochemistry, Example Institute, New Delhi, India.")
d.add_paragraph("ABSTRACT")
d.add_paragraph("Background: Tumour protein p53 (TP53) is mutated in many cancers. Methods: MCF-7 and MDA-MB-231 cells were treated and analysed by qPCR and Western blot. Results: TP53 expression rose 2.1-fold (95% CI 1.4–2.8, p < 0.01). Conclusion: The compound modulates TP53.")
d.add_paragraph("KEYWORDS: TP53; BRCA1; breast cancer; apoptosis")
d.add_paragraph("INTRODUCTION")
d.add_paragraph("Breast cancer is the most common cancer in women worldwide [1]. Mutations in TP53 and BRCA1 drive tumour progression [2,3]. Interleukin-6 (IL-6) and tumour necrosis factor alpha (TNF-α) promote inflammation (Fig. 1A).")
d.add_paragraph("MATERIALS AND METHODS")
d.add_paragraph("Cell lines MCF-7 and MDA-MB-231 were obtained from ATCC and grown in DMEM with 10% FBS at 37 °C and 5% CO2. Cells were washed with PBS and lysed in RIPA buffer containing NaCl, MgCl2 and EDTA. Protein was quantified by BCA assay and separated by SDS-PAGE.")
d.add_paragraph("RNA was extracted with TRIzol and cDNA synthesised. qPCR was performed for TP53, BRCA1, BAX and BCL2 using GAPDH as reference (Table 1). Data are mean ± SD of three experiments; groups were compared by one-way ANOVA in SPSS v26 and GraphPad Prism.")
d.add_paragraph("The study was conducted between June 2023 and March 2025. No human participants or animals were involved; ethical approval was therefore not required.")
d.add_paragraph("RESULTS")
d.add_paragraph("Treatment increased BAX and decreased BCL2 (Figure 1B, Figure 2). Apoptosis rose in both cell lines (Figures 2–3; Supplementary Figure S1). Results are summarised in Table 1 and Table 2.")
d.add_paragraph("DISCUSSION")
d.add_paragraph("Our findings agree with earlier work on TP53 activation [4]. The IL-6 pathway may be involved, as reported in the USA and UK cohorts [5].")
d.add_paragraph("CONCLUSIONS")
d.add_paragraph("The compound shows promise.")
d.add_paragraph("FUNDING")
d.add_paragraph("This research received no external funding.")
d.add_paragraph("CONFLICTS OF INTEREST")
d.add_paragraph("The authors declare no conflict of interest.")
d.add_paragraph("AUTHOR CONTRIBUTIONS")
d.add_paragraph("Conceptualization, R.S.; investigation, K.I.; writing, R.S. and K.I.")
d.add_paragraph("DATA AVAILABILITY STATEMENT")
d.add_paragraph("Data are available from the corresponding author on request.")
d.add_paragraph("ACKNOWLEDGMENTS")
d.add_paragraph("We thank the CSIR for infrastructure support.")
d.add_paragraph("Figure 1. Expression of TP53. (A) qPCR. (B) Western blot.")
d.add_paragraph("Figure 2. BAX/BCL2 ratio.")
d.add_paragraph("Figure 3. Apoptosis by flow cytometry.")
d.add_paragraph("Table 1. Primer sequences.")
d.add_paragraph("Table 2. Summary of IC50 values.")
t = d.add_table(rows=2, cols=2); t.cell(0,0).text = "Gene"; t.cell(0,1).text = "Primer"
t = d.add_table(rows=2, cols=2); t.cell(0,0).text = "Cell line"; t.cell(0,1).text = "IC50"
d.add_paragraph("REFERENCES")
for r in ["1. Sung H, Ferlay J. Global cancer statistics 2020. CA Cancer J Clin. 2021;71:209-49. doi:10.3322/caac.21660",
          "2. Olivier M, Hollstein M. TP53 mutations in human cancers. CSH Perspect Biol. 2010;2:a001008.",
          "3. Miki Y. A strong candidate for BRCA1. Science. 1994;266:66-71.",
          "4. Vousden KH. p53 in health and disease. Nat Rev Mol Cell Biol. 2007;8:275-83.",
          "5. Hunter T. IL-6 in cancer. JAMA Oncol. 2019;5:1-8."]:
    d.add_paragraph(r)
d.save("tests/realistic_paper.docx")
print("saved")
